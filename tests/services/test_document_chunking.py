from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image

from baozhi_rag.services.document_chunking import DocumentChunkService
from baozhi_rag.services.term_matching import TermMatchResult


class _StubTermMatcher:
    """返回空领域词结果，避免测试依赖真实词典内容。"""

    def extract_terms(self, text: str) -> TermMatchResult:
        """返回固定的空匹配结果。"""
        return TermMatchResult(merged_terms=[])


def _build_chunk_service(tmp_path: Path, *, chunk_size: int = 200) -> DocumentChunkService:
    """构造测试用的文档切块服务。"""
    return DocumentChunkService(
        chunk_size=chunk_size,
        chunk_overlap=20,
        convert_temp_dir=tmp_path / "converted",
        term_matcher=_StubTermMatcher(),
    )


def test_chunk_docx_keeps_paragraph_comment_text(tmp_path: Path) -> None:
    """段落批注应被追加到可检索的 chunk 内容中。"""
    file_path = tmp_path / "paragraph-comment.docx"
    document = Document()
    paragraph = document.add_paragraph()
    first_run = paragraph.add_run("Main clause")
    second_run = paragraph.add_run(" appendix")
    document.add_comment([first_run, second_run], text="review note", author="tester")
    document.save(file_path)

    service = _build_chunk_service(tmp_path)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/paragraph-comment.docx",
        file_id="file-1",
    )

    assert len(chunks) == 1
    assert "Main clause appendix" in chunks[0].content
    assert "review note" in chunks[0].content


def test_chunk_docx_keeps_table_comment_text_when_large_table_is_split(tmp_path: Path) -> None:
    """超长表格拆分后，每个分片都应保留表格批注文本与表头。"""
    file_path = tmp_path / "table-comment.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"

    for index in range(1, 7):
        row = table.add_row()
        row.cells[0].text = f"Row {index} value"
        row.cells[1].text = f"Description {index} with extra words"

    comment_paragraph = table.cell(1, 0).paragraphs[0]
    comment_run = comment_paragraph.runs[0]
    document.add_comment(comment_run, text="table review", author="tester")
    document.save(file_path)

    service = _build_chunk_service(tmp_path, chunk_size=120)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/table-comment.docx",
        file_id="file-2",
    )

    assert len(chunks) >= 2
    assert all("Header A" in chunk.content for chunk in chunks)
    assert all("table review" in chunk.content for chunk in chunks)


def test_chunk_docx_extracts_embedded_images(tmp_path: Path) -> None:
    """段落内嵌图片应被抽取为 chunk 图片资产。"""
    file_path = tmp_path / "paragraph-image.docx"
    document = Document()
    paragraph = document.add_paragraph("Image paragraph")

    image_buffer = BytesIO()
    Image.new("RGB", (16, 16), color=(255, 0, 0)).save(image_buffer, format="PNG")
    image_buffer.seek(0)
    paragraph.add_run().add_picture(image_buffer)
    document.save(file_path)

    service = _build_chunk_service(tmp_path)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/paragraph-image.docx",
        file_id="file-3",
    )

    assert len(chunks) == 1
    assert len(chunks[0].image_assets) == 1
    assert chunks[0].image_assets[0].content_type == "image/png"
    assert chunks[0].image_assets[0].image_bytes is not None


def test_chunk_docx_keeps_heading_path_and_section_title(tmp_path: Path) -> None:
    """正文 chunk 应保留标题路径与末级标题。"""
    file_path = tmp_path / "heading-path.docx"
    document = Document()
    document.add_heading("系统概览", level=1)
    document.add_paragraph("这里是概览段落。")
    document.save(file_path)

    service = _build_chunk_service(tmp_path)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/heading-path.docx",
        file_id="file-4",
    )

    assert len(chunks) >= 1
    assert chunks[0].heading_path == ["系统概览"]
    assert chunks[0].section_title == "系统概览"
    assert chunks[0].content_type == "paragraph"


def test_chunk_docx_splits_buffer_when_heading_changes(tmp_path: Path) -> None:
    """标题上下文变化时，应避免跨章节把多个段落合并到同一 chunk。"""
    file_path = tmp_path / "heading-buffer.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("第一章内容。")
    document.add_heading("第二章", level=1)
    document.add_paragraph("第二章内容。")
    document.save(file_path)

    service = _build_chunk_service(tmp_path, chunk_size=500)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/heading-buffer.docx",
        file_id="file-5",
    )

    assert len(chunks) >= 2
    assert chunks[0].section_title == "第一章"
    assert chunks[1].section_title == "第二章"


def test_chunk_docx_marks_table_chunk_content_type(tmp_path: Path) -> None:
    """表格 chunk 应显式标记为 table。"""
    file_path = tmp_path / "table-type.docx"
    document = Document()
    document.add_heading("配置表", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "键"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "timeout"
    table.cell(1, 1).text = "30"
    document.save(file_path)

    service = _build_chunk_service(tmp_path, chunk_size=500)
    chunks = service.chunk_docx(
        file_path=file_path,
        source_filename=file_path.name,
        storage_key="stage/table-type.docx",
        file_id="file-6",
    )

    table_chunks = [chunk for chunk in chunks if chunk.content_type == "table"]
    assert table_chunks
    assert table_chunks[0].heading_path == ["配置表"]
    assert table_chunks[0].section_title == "配置表"
